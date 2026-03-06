import os
import re
from pypdf import PdfReader
from docx import Document
import openpyxl

DOCS_PATH = "docs"


def _normalizar(texto: str) -> str:
    texto = texto.lower()
    texto = re.sub(r"[^\w\sáàâãéêíóôõúç]", " ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto


def ler_txt(caminho):
    with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def ler_pdf(caminho):
    texto = []
    reader = PdfReader(caminho)
    for page in reader.pages:
        try:
            texto.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(texto)


def ler_docx(caminho):
    doc = Document(caminho)
    return "\n".join([p.text for p in doc.paragraphs])


def ler_xlsx(caminho):
    wb = openpyxl.load_workbook(caminho)
    texto = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            linha = " ".join([str(c) for c in row if c is not None])
            if linha.strip():
                texto.append(linha.strip())
    return "\n".join(texto)


def ler_documento(caminho):
    nome = caminho.lower()
    if nome.endswith(".txt"):
        return ler_txt(caminho)
    if nome.endswith(".pdf"):
        return ler_pdf(caminho)
    if nome.endswith(".docx"):
        return ler_docx(caminho)
    if nome.endswith(".xlsx"):
        return ler_xlsx(caminho)
    return ""


def buscar_docs(pergunta: str, limite=6):
    """
    Busca por relevância: quebra a pergunta em palavras e pontua linhas/arquivos.
    Retorna uma lista de trechos (arquivo + linha) mais prováveis de responder.
    """
    if not os.path.exists(DOCS_PATH):
        return []

    q_norm = _normalizar(pergunta)
    palavras = [p for p in q_norm.split() if len(p) >= 3]

    candidatos = []

    for arquivo in os.listdir(DOCS_PATH):
        caminho = os.path.join(DOCS_PATH, arquivo)
        try:
            conteudo = ler_documento(caminho)
        except Exception:
            continue

        if not conteudo:
            continue

        for linha in conteudo.splitlines():
            ln = linha.strip()
            if not ln:
                continue

            ln_norm = _normalizar(ln)
            score = sum(1 for p in palavras if p in ln_norm)

            # bônus se a linha contém "multa", "prazo", etc.
            if "multa" in ln_norm:
                score += 2 
                if "%" in ln:
                    score += 2
            if "prazo" in ln_norm:
                score += 1

            if score > 0:
                candidatos.append((score, arquivo, ln))

    candidatos.sort(key=lambda x: x[0], reverse=True)

    vistos = set()
    saida = []
    for score, arquivo, ln in candidatos:
        key = (arquivo, ln)
        if key in vistos:
            continue
        vistos.add(key)
        saida.append(f"{arquivo}: {ln}")
        if len(saida) >= limite:
            break

    return saida